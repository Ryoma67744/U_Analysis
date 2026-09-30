"""★ ver77.0: 手法取り違え、出力中の置換、再入力feature汚染を実writerで検証する。"""
from collections import OrderedDict
from copy import deepcopy
import json
import os
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from app.callbacks import interactive_data_export as ex
from app.services import result_snapshot as rs
from app.services.parquet_column_roles import (ROLE_KEY, MANIFEST_KEY, export_metadata,
                                              feature_columns, read_column_roles)


def _descriptor(path, method, *, embedding="umap", clusters="computed"):
    return {"descriptor_version": 1, "result_id": method,
            "source": {"path": str(path), "sha256": rs.file_signature(path)},
            "classification": {"method": method, "state": "recorded"},
            "embedding": {"kind": embedding}, "clusters": {"kind": clusters},
            "capabilities": {"spatial": True, "deg": clusters == "computed",
                             "umap": embedding == "umap", "pc": embedding == "pca2d"}}


@pytest.fixture
def results(tmp_path, monkeypatch):
    folder = tmp_path / "raw"; folder.mkdir()
    root = tmp_path / "result"; root.mkdir()
    raw = pd.DataFrame({"id": [3, 1, 2], "x": [3., 1., 2.], "y": 0.,
                        "500.000000": [30., 10., 20.], "annotation": "sample"})
    raw.to_parquet(folder / "sample.parquet", index=False)
    frames, rmap, descriptors = OrderedDict(), {}, {}
    for method, offset, labels in [("Harmony", 10., ["0", "1", "2"]),
                                    ("RPCA", 80., ["7", "8", "9"])]:
        path = root / f"{method}.rds"; path.write_bytes(method.encode())
        rmap[method] = str(path)
        frames[method] = pd.DataFrame({"CellID": ["a", "b", "c"], "Sample": "sample",
            "SpatialX": [1., 2., 3.], "SpatialY": 0., "Cluster": labels,
            "UMAP_1": np.arange(3.) + offset, "UMAP_2": np.arange(3.) - offset,
            "TotalCount": [101., 102., 103.]})
        descriptors[method] = _descriptor(path, method)
    calls = []
    def extract(path, force_verify=False):
        method = Path(path).stem; calls.append((method, force_verify))
        return {"plot_data": frames[method], "meta": {}, "result_descriptor": descriptors[method]}
    monkeypatch.setattr(ex._bridge, "extract_data", extract)
    monkeypatch.setattr(ex, "_interactive_data", {"method": "RPCA", "plot_data": frames["RPCA"]})
    monkeypatch.setattr(ex, "_build_region_lookup", lambda *_a: (None, []))
    return SimpleNamespace(folder=folder, root=root, frames=frames, rmap=rmap,
                           descriptors=descriptors, calls=calls, raw=raw)


def _run(results, tmp_path, methods, *, snapshot=None):
    return ex._do_export(str(results.folder), "TIMS", "parquet", results.rmap, "RPCA",
        str(results.root), None, None, results.rmap["RPCA"], selected_methods=methods,
        options={"categories": ["id", "coords", "intensity", "section", "cluster", "umap", "quality"]},
        out_dir=tmp_path / "downloads", snapshot=snapshot)


def test_display_rpca_export_harmony_has_one_source_for_clusters_and_coordinates(results, tmp_path):
    path, _, message = _run(results, tmp_path, ["Harmony"])
    assert path, message
    out = pd.read_parquet(path)
    assert out["UMAP cluster"].tolist() == ["2", "0", "1"]
    assert out["UMAP_1"].tolist() == [12., 10., 11.]
    assert out["500.000000"].tolist() == [30., 10., 20.]
    assert results.calls == [("Harmony", True)]
    schema = pq.read_schema(path)
    roles = read_column_roles(schema)
    assert roles["UMAP_1"]["method"] == "Harmony"
    assert feature_columns(schema.names, roles) == ["500.000000"]
    manifest = json.loads(schema.metadata[MANIFEST_KEY])
    assert manifest["conditions"]["extra"]["method_results"][0]["requested_method"] == "Harmony"
    assert str(results.root) not in schema.metadata[MANIFEST_KEY].decode("utf-8")
    assert manifest["input_files"][0]["path"] == "sample.parquet"
    assert manifest["input_files"][0]["sha256"] == rs.file_signature(results.folder / "sample.parquet")


def test_multiple_method_subset_nulls_are_consistent_for_cluster_and_embedding(results, tmp_path):
    results.frames["Harmony"] = results.frames["Harmony"].iloc[:2].copy()
    results.frames["RPCA"] = results.frames["RPCA"].iloc[1:].copy()
    path, _, message = _run(results, tmp_path, ["Harmony", "RPCA"])
    assert path, message
    out = pd.read_parquet(path)
    assert "UMAP_1" not in out
    assert out["Harmony"].tolist() == [None, "0", "1"]
    assert out["RPCA"].tolist() == ["9", None, "8"]
    assert pd.isna(out.loc[0, "Harmony__UMAP_1"])
    assert pd.isna(out.loc[1, "RPCA__UMAP_1"])
    assert out.loc[2, "Harmony__UMAP_1"] == 11.
    assert out.loc[2, "RPCA__UMAP_1"] == 81.
    assert feature_columns(out.columns, read_column_roles(pq.read_schema(path))) == ["500.000000"]


def test_pc_fallback_is_not_exported_as_umap(results, tmp_path):
    results.descriptors["Harmony"] = _descriptor(results.rmap["Harmony"], "PCA", embedding="pca2d", clusters="inherited")
    path, _, message = _run(results, tmp_path, ["Harmony"])
    assert path, message
    out = pd.read_parquet(path)
    assert "PC_1" in out and "UMAP_1" not in out
    roles = read_column_roles(pq.read_schema(path))
    assert roles["PC_1"]["method"] == "PCA_projection"
    assert roles["PC_1"]["embedding_kind"] == "pca2d"


def test_request_rejects_same_size_same_mtime_replacement(results):
    request = rs.build_export_request(results.rmap, ["Harmony"])
    path = Path(results.rmap["Harmony"]); stat = path.stat()
    path.write_bytes(b"changed")
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns))
    with pytest.raises(ValueError, match="変更"):
        rs.load_method_snapshots(request, ex._bridge)
    assert not results.calls


@pytest.mark.parametrize("change", ["replace", "add"])
def test_request_detects_embedding_sidecar_change_even_when_primary_is_unchanged(results, change):
    sidecar = results.root / "UMAP_Harmony.rds"
    if change == "replace":
        sidecar.write_bytes(b"embedding-v1")
    request = rs.build_export_request(results.rmap, ["Harmony"])
    primary_hash = rs.file_signature(results.rmap["Harmony"])
    sidecar.write_bytes(b"embedding-v2")
    assert rs.file_signature(results.rmap["Harmony"]) == primary_hash
    with pytest.raises(ValueError, match="補助ファイル"):
        rs.load_method_snapshots(request, ex._bridge)
    with pytest.raises(ValueError, match="補助ファイル"):
        rs.verify_request(request)


def test_cheap_acceptance_stamp_rejects_replacement_before_worker(results):
    states = rs.capture_source_states(results.rmap, ["Harmony"])
    Path(results.rmap["Harmony"]).write_bytes(b"updated source")
    with pytest.raises(ValueError, match="受付後"):
        rs.build_export_request(results.rmap, ["Harmony"], expected_states=states)


def test_bridge_sidecar_descriptor_must_match_pinned_sidecar(results):
    sidecar = results.root / "UMAP_Harmony.rds"; sidecar.write_bytes(b"embedding")
    results.descriptors["Harmony"]["embedding"].update(location="sidecar",
        source_path=str(sidecar), source_sha256="stale")
    request = rs.build_export_request(results.rmap, ["Harmony"])
    with pytest.raises(ValueError, match="補助ファイル"):
        rs.load_method_snapshots(request, ex._bridge)


def test_source_identity_must_also_preserve_spatial_coordinates():
    from app.services.section_group_metadata import SourceClusterLookup
    lookup = SourceClusterLookup(); lookup.by_source[("f", "1")] = "0"
    frame = pd.DataFrame({"source_file_id": ["f"], "source_pixel_id": ["1"],
                          "SpatialX": [4.], "SpatialY": [8.]})
    audit = rs.ExportJoinAudit({"PCA": lookup}, reference_frame=frame)
    with pytest.raises(ValueError, match="空間座標"):
        audit.consume([None], [("f", "1")], [(9., 8.)])


@pytest.mark.parametrize("coordinates", [[1., 1.], [1.000001, 1.000002]])
def test_legacy_lookup_rejects_duplicate_pixels_before_dictionary_collapse(coordinates):
    frame = pd.DataFrame({"CellID": ["a", "b"], "Sample": "sample",
                          "SpatialX": coordinates, "SpatialY": 0., "Cluster": ["1", "2"]})
    with pytest.raises(ValueError, match="複数の画素"):
        ex._build_cluster_lookup(frame)


def test_complete_source_identity_preserves_distinct_pixels_at_same_coordinates():
    frame = pd.DataFrame({"CellID": ["a", "b"], "Sample": "sample",
                          "SpatialX": 1., "SpatialY": 0., "Cluster": ["1", "2"],
                          "source_file_id": "f", "source_pixel_id": ["001", "002"]})
    lookup = ex._build_cluster_lookup(frame)
    assert lookup.by_source == {("f", "001"): "1", ("f", "002"): "2"}
    audit = rs.ExportJoinAudit({"PCA": lookup}, reference_frame=frame)
    audit.consume([("sample", 1., 0.), ("sample", 1., 0.)], [("f", "001"), ("f", "002")])
    audit.validate()
    with pytest.raises(ValueError, match="複数の画素"):
        ex._build_cluster_lookup(frame.assign(source_pixel_id=["001", None]))


def test_api_acceptance_freezes_mapping_and_does_not_hash_large_rds(results, monkeypatch):
    import flask
    import threading
    from app.services import gpt_api as api, export_progress as progress
    pending = []
    class DeferredThread:
        def __init__(self, target, args, daemon): self.target, self.args = target, args
        def start(self): pending.append(self)
    monkeypatch.setattr(api.threading, "Thread", DeferredThread)
    gate = threading.BoundedSemaphore(1)
    monkeypatch.setattr(api, "_GPT_EXPORT_SEM", gate)
    source = {"project": {"id": "p"}, "sub": {"id": "s"}, "rds_map": deepcopy(results.rmap),
              "result_dir": str(results.root), "data_folder": str(results.folder), "ms_instrument": "TIMS"}
    monkeypatch.setattr(api, "_resolve_sub", lambda *_a: source)
    monkeypatch.setattr("app.config.GPT_API_KEY", "test-export-key")
    monkeypatch.setattr(rs, "file_signature", lambda *_a: pytest.fail("受付スレッドでRDSをhashした"))
    app = flask.Flask(__name__); api.register_gpt_api(app)
    response = app.test_client().post("/api/gpt/projects/p/sub/s/exports/interactive?methods=Harmony",
                                     headers={"X-API-Key": "test-export-key"})
    try:
        assert response.status_code == 200, response.get_json()
        assert len(pending) == 1
        captured = pending[0].args[1]
        source["rds_map"].clear()
        assert captured["rds_map"]["Harmony"] == results.rmap["Harmony"]
        assert captured["_export_source_states"]
        assert "_export_manifest" in captured
    finally:
        if pending:
            progress.pop_job(pending[0].args[0])
        gate.release()


def test_snapshot_does_not_alias_extracted_dataframe(results):
    request = rs.build_export_request(results.rmap, ["Harmony"])
    snapshot = rs.load_method_snapshots(request, ex._bridge)["Harmony"]
    results.frames["Harmony"].loc[0, "UMAP_1"] = 1234.
    frame = snapshot.frame(); frame.loc[0, "UMAP_1"] = 5678.
    assert snapshot.frame().loc[0, "UMAP_1"] == 10.


@pytest.mark.parametrize("failure", ["missing", "duplicate"])
def test_unexpected_join_failure_has_no_published_file(results, tmp_path, failure):
    raw = results.raw.iloc[:2].copy() if failure == "missing" else pd.concat([results.raw, results.raw.iloc[[0]]])
    raw.to_parquet(results.folder / "sample.parquet", index=False)
    path, _, message = _run(results, tmp_path, ["Harmony"])
    assert path is None and ("結合" in message or "複数" in message)
    assert not list((tmp_path / "downloads").glob("*.parquet"))


def test_source_replacement_during_serialization_is_not_published(results, tmp_path, monkeypatch):
    original = pd.DataFrame.to_parquet
    def changing(frame, path, *args, **kwargs):
        original(frame, path, *args, **kwargs)
        Path(results.rmap["Harmony"]).write_bytes(b"new RDS")
    monkeypatch.setattr(pd.DataFrame, "to_parquet", changing)
    path, _, message = _run(results, tmp_path, ["Harmony"])
    assert path is None and "変更" in message
    assert not list((tmp_path / "downloads").glob("*.parquet"))


def test_raw_intensity_replacement_during_serialization_is_not_published(results, tmp_path, monkeypatch):
    original = pd.DataFrame.to_parquet
    def changing(frame, path, *args, **kwargs):
        original(frame, path, *args, **kwargs)
        original(results.raw.assign(**{"500.000000": [1000., 2000., 3000.]}),
                 results.folder / "sample.parquet", index=False)
    monkeypatch.setattr(pd.DataFrame, "to_parquet", changing)
    path, _, message = _run(results, tmp_path, ["Harmony"])
    assert path is None and "入力ファイルが変更" in message
    assert not list((tmp_path / "downloads").glob("*.parquet"))


def test_known_analysis_input_hash_must_match(results):
    path = results.folder / "sample.parquet"
    manifest = {"files": [{"path": str(path), "validation": {"parquet": {"sha256": "stale"}}}]}
    with pytest.raises(ValueError, match="解析時の入力"):
        rs.pin_input_files([path], manifest)


def test_api_driver_uses_selected_result_and_same_parquet_metadata(results, tmp_path):
    path, _, message = ex.build_interactive_export_for_project(
        str(results.folder), "TIMS", "parquet", results.rmap, str(results.root), None, None,
        selected_methods=["Harmony"], out_dir=tmp_path / "api")
    assert path, message
    out = pd.read_parquet(path)
    assert out["UMAP cluster"].tolist() == ["2", "0", "1"]
    roles = read_column_roles(pq.read_schema(path))
    assert roles["UMAP cluster"]["method"] == "Harmony"
    assert feature_columns(out.columns, roles) == ["500.000000"]


def test_desi_validates_sources_after_excel_serialization(results, tmp_path, monkeypatch):
    from openpyxl.workbook.workbook import Workbook
    folder = tmp_path / "desi"; folder.mkdir()
    source = folder / "sample.txt"
    source.write_text("header\tx\ty\n" * 5 + "1\t1\t0\n2\t2\t0\n3\t3\t0\n", encoding="utf-8")
    original = Workbook.save
    def save(book, filename):
        original(book, filename)
        source.write_text("changed input", encoding="utf-8")
    monkeypatch.setattr(Workbook, "save", save)
    request = rs.build_export_request(results.rmap, ["Harmony"])
    with pytest.raises(ValueError, match="入力ファイルが変更"):
        ex._export_desi(str(folder), OrderedDict(Harmony=ex._build_cluster_lookup(results.frames["Harmony"])),
                        out_dir=tmp_path / "desi_output", export_request=request)
    assert not list((tmp_path / "desi_output").glob("*.xlsx"))


def test_missing_pca_export_does_not_generate_a_projection(results, monkeypatch):
    results.rmap["PCA"] = str(results.root / "missing_pca.rds")
    monkeypatch.setattr(ex._bridge, "derive_uncorrected_pca", lambda *_a: pytest.fail("出力時にPCA生成した"))
    with pytest.raises(ValueError, match="解析結果がありません"):
        ex._selection_snapshot(results.rmap["Harmony"], results.rmap, "Harmony", ["PCA"], str(results.root))


def test_declared_roles_drive_python_input_and_reject_malformed_metadata(tmp_path):
    from app.services.data_manager import _read_tims_raw
    frame = pd.DataFrame({"id": [1], "x": [0.], "y": [0.], "500.000000": [7.],
                          "custom_method_1": [80.], "RPCA__UMAP_1": [9.]})
    metadata = export_metadata(list(frame.columns), overrides={"custom_method_1": {"role": "cluster"}})
    table = pa.Table.from_pandas(frame, preserve_index=False).replace_schema_metadata(metadata)
    path = tmp_path / "analysis.parquet"; pq.write_table(table, path)
    result = _read_tims_raw(tmp_path)
    # readerの数値特徴軸へ埋込やクラスタが入り込まない。
    assert list(result.columns) == ["500.000000"] and result.iloc[0, 0] == 7.
    broken = table.replace_schema_metadata({ROLE_KEY: b'{"schema_version":1,"columns":[]}'})
    pq.write_table(broken, path)
    with pytest.raises(ValueError, match="列役割"):
        _read_tims_raw(tmp_path)


def test_legacy_coordinate_names_are_never_features():
    names = ["id", "x", "y", "500.1", "UMAP_1", "PCA__PC_2", "RPCA__UMAP_2", "TotalCount"]
    assert feature_columns(names) == ["500.1"]


@pytest.mark.parametrize("mode", ["legacy_single", "selected_single", "multiple"])
def test_real_pptx_all_branches_use_fixed_method_result_and_explain_projection(results, tmp_path, monkeypatch, mode):
    import base64
    from io import BytesIO
    import sys
    from types import ModuleType
    from pptx import Presentation
    from app.callbacks import interactive_pptx as ppt
    from dash import no_update
    monkeypatch.setitem(sys.modules, "kaleido", ModuleType("kaleido"))
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")
    monkeypatch.setattr(ppt, "_fig_to_png_bytes", lambda *_a, **_k: png)
    monkeypatch.setattr(ppt._bridge, "ensure_expression_matrix", lambda *_a, **_k: None)
    deg_calls = []
    monkeypatch.setattr(ppt, "_load_deg_results", lambda *args, **kwargs: deg_calls.append((args, kwargs)))
    results.descriptors["Harmony"] = _descriptor(results.rmap["Harmony"], "PCA", embedding="pca2d", clusters="inherited")
    results.frames["Harmony"]["Cluster"] = "0"
    results.frames["RPCA"]["Cluster"] = "7"
    captured = []
    spatial_inputs = []
    original_spatial = ppt._create_single_spatial_fig
    def spatial(frame, *args, **kwargs):
        spatial_inputs.append(frame[["Cluster", "UMAP_1", "UMAP_2"]].copy())
        return original_spatial(frame, *args, **kwargs)
    monkeypatch.setattr(ppt, "_create_single_spatial_fig", spatial)
    original = ppt._build_pptx
    def build(*args, **kwargs):
        captured.append((args[4], kwargs["df"].copy(), deepcopy(kwargs["conditions"]), kwargs["result_descriptor"]))
        return original(*args, **kwargs)
    monkeypatch.setattr(ppt, "_build_pptx", build)
    current = "RPCA" if mode != "legacy_single" else "Harmony"
    rmap = None if mode == "legacy_single" else results.rmap
    methods = [] if mode == "legacy_single" else ["Harmony", "RPCA"] if mode == "multiple" else ["Harmony"]
    # 故意に別結果のcache_dirを渡しても旧fastpathで読まない。
    cache = tmp_path / "stale_cache"; cache.mkdir()
    results.frames["RPCA"].to_parquet(cache / "plot_data.parquet")
    download, message = ppt.cb_export_report(lambda *_a: None, 1,
        {"data": [], "layout": {}}, None, results.rmap[current], [], None, None,
        None, None, None, {}, None, None, 5, str(cache), None, rmap, str(results.root),
        current, methods, False, {}, None)
    assert download is not no_update, message
    document = Presentation(BytesIO(base64.b64decode(download["content"])))
    assert len(document.slides) >= 5
    text = "\n".join(shape.text for slide in document.slides for shape in slide.shapes if shape.has_text_frame)
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in document.slides)
    assert "PC1" in text and "既存クラスタ" in text
    assert '"result_descriptor"' in notes and '"kind": "inherited"' in notes
    assert str(results.root) not in notes
    expected = {"Harmony", "RPCA"} if mode == "multiple" else {"Harmony"}
    assert {Path(path).stem for path, *_rest in captured} == expected
    assert spatial_inputs
    for frame in spatial_inputs:
        method = "Harmony" if set(frame["Cluster"]) == {"0"} else "RPCA"
        assert method in expected
        assert set(frame["UMAP_1"]).issubset(set(results.frames[method]["UMAP_1"]))
    for path, frame, conditions, descriptor in captured:
        method = Path(path).stem
        np.testing.assert_array_equal(frame[["UMAP_1", "UMAP_2"]], results.frames[method][["UMAP_1", "UMAP_2"]])
        assert frame["Cluster"].tolist() == results.frames[method]["Cluster"].tolist()
        assert conditions["extra"]["method_result"]["source_sha256"] == rs.file_signature(path)
        assert conditions["extra"]["method_result"]["requested_method"] == method
    assert len(deg_calls) == (1 if mode == "multiple" else 0)
    if deg_calls:
        args, kwargs = deg_calls[0]
        assert args == (results.root, "RPCA")
        assert kwargs["strict_method"] is True
        assert kwargs["result_descriptor"] == results.descriptors["RPCA"]


@pytest.mark.parametrize("ownership", ["actual_method", "unnamed", "other_method", "inherited", "unknown"])
def test_pptx_deg_uses_verified_method_and_source_root(results, tmp_path, monkeypatch, ownership):
    import base64
    import sys
    from io import BytesIO
    from types import ModuleType
    from pptx import Presentation
    from dash import no_update
    from app.callbacks import interactive_pptx as ppt
    monkeypatch.setitem(sys.modules, "kaleido", ModuleType("kaleido"))
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")
    monkeypatch.setattr(ppt, "_fig_to_png_bytes", lambda *_a, **_k: png)
    monkeypatch.setattr(ppt._bridge, "ensure_expression_matrix", lambda *_a, **_k: None)
    results.descriptors["Harmony"] = _descriptor(results.rmap["Harmony"], "PCA",
        clusters="inherited" if ownership == "inherited" else "computed")
    if ownership == "unknown":
        results.descriptors["Harmony"]["classification"] = {"method": None, "state": "unknown"}
    unrelated = tmp_path / "unrelated_result"; unrelated.mkdir()
    def markers(root, method, gene):
        folder = root / method if method else root
        folder.mkdir(exist_ok=True)
        pd.DataFrame({"gene": [gene], "cluster": ["0"], "avg_log2FC": [2.],
                      "p_val_adj": [0.001]}).to_csv(folder / "markers_annotated.csv", index=False)
    markers(unrelated, "PCA", "wrong-root")
    markers(results.root, "Harmony", "wrong-method")
    if ownership != "other_method":
        markers(results.root, "" if ownership == "unnamed" else "PCA", "correct-pca")
    captured = []
    def build(*args, **kwargs):
        captured.append(kwargs)
        return kwargs.get("progress_offset", 0)
    monkeypatch.setattr(ppt, "_build_pptx", build)
    download, message = ppt.cb_export_report(lambda *_a: None, 1,
        {"data": [], "layout": {}}, None, results.rmap["Harmony"], [], None, None,
        None, None, None, {}, None, None, 5, None, None, results.rmap, str(unrelated),
        "Harmony", ["Harmony"], False, {}, None)
    assert download is not no_update, message
    document = Presentation(BytesIO(base64.b64decode(download["content"])))
    assert document.slides and len(captured) == 1
    records = captured[0]["deg_data"]
    if ownership == "actual_method":
        assert [record["gene"] for record in records] == ["correct-pca"]
    else:
        assert records is None
