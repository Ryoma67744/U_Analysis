"""★ ver75.1: 実writerで対象選択と強度・座標・集計の対応を確認する。"""
from collections import OrderedDict
from copy import deepcopy
import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from app.callbacks import interactive_data_export as exp
from app.services.export_selection import build_catalog, resolve_selection
from app.services.export_mzlist import SHEET_NAME
from app.services.section_metadata import build_section_manifest

SCOPE = "export-selection-io-result"


def _choose(frames, groups):
    catalog = build_catalog(frames, SCOPE)
    request = {"mode": "selected", "scope": SCOPE, "signature": catalog["signature"],
               "ids": [unit["id"] for unit in catalog["units"] if unit["group"] in groups]}
    return resolve_selection(frames, request, SCOPE)


def _lookups(frames):
    return OrderedDict((method, exp._build_cluster_lookup(frame)) for method, frame in frames.items())


def _read(path, fmt, sheet="Data"):
    if fmt == "parquet":
        return pd.read_parquet(path)
    if fmt == "xlsx":
        return pd.read_excel(path, sheet_name=sheet)
    return pd.read_csv(path)


@pytest.fixture
def modern(tmp_path):
    # 同じbasename/ROI/座標でも元ファイルIDで区別できなければ科学結果が混ざる。
    entries, raws = [], {}
    for index, group in enumerate(("Ctrl", "KO1", "KO2")):
        folder = tmp_path / group
        folder.mkdir()
        source = folder / "same.parquet"
        raw = pd.DataFrame({"id": [2, 1, 99], "x": [20., 10., 99.], "y": [0., 0., 0.],
                            "500.100000": [index * 10 + 4., index * 10 + 2., 900.],
                            "600.200000": [index * 100 + 8., index * 100 + 4., 1800.],
                            "annotation": ["ROI"] * 3})
        raw.to_parquet(source, index=False)
        raws[str(source)] = raw
        entries.append({"path": str(source), "available_rois": ["ROI"], "roi_role": "section"})
    manifest = build_section_manifest(entries)
    records = []
    for index, (entry, group) in enumerate(zip(manifest["files"], ("Ctrl", "KO1", "KO2"))):
        section = entry["sections"][0]
        section.update(subject_id=f"Mouse{index}", group=group, metadata_confirmed=True)
        for pixel in (1, 2):
            records.append({"CellID": f"{index}:{pixel}", "Sample": "display name shared across files",
                            "source_file_id": entry["file_id"], "source_pixel_id": str(pixel),
                            "section_id": section["section_id"], "section_display_name": f"Slice{index}",
                            "subject_id": f"Mouse{index}", "group": group,
                            "integration_unit_id": section["section_id"],
                            "SpatialX": pixel * 10., "SpatialY": 0., "Cluster": str(index + pixel),
                            "UMAP_1": index + pixel * .125, "UMAP_2": index - pixel * .25})
    frame = pd.DataFrame(records)
    rds = tmp_path / "result" / "RDS_Files" / "PCA.rds"
    rds.parent.mkdir(parents=True)
    rds.write_bytes(b"immutable RDS fixture")
    (rds.parent.parent / "section_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return {"rds": rds, "manifest": manifest, "frames": OrderedDict(PCA=frame), "raws": raws,
            "folder": str(Path(entries[0]["path"]).parent)}


@pytest.mark.parametrize("fmt", ["csv", "parquet", "xlsx"])
@pytest.mark.parametrize("exclude_unused", [False, True])
def test_modern_group_selection_writes_only_chosen_pixels_without_changing_values(modern, tmp_path, fmt, exclude_unused):
    frames = modern["frames"]
    before = frames["PCA"].copy(deep=True)
    files_before = {path: Path(path).read_bytes() for path in modern["raws"]}
    opts = {"categories": ["id", "coords", "intensity", "section", "umap", "cluster"]}
    selected = _choose(frames, {"Ctrl", "KO2"})
    conditions = {"extra": {}}
    path, name = exp._export_tims(modern["folder"], _lookups(frames), fmt,
        selection=selected, metadata_rds=modern["rds"], metadata_plot=before,
        exclude_unused=exclude_unused, options=opts, extra_lookups=exp._build_extra_lookups(before, opts),
        conditions=conditions, out_dir=tmp_path)
    out = _read(path, fmt)
    assert name.endswith(f"_selected.{fmt}")
    assert out["group"].tolist() == ["Ctrl", "Ctrl", "KO2", "KO2"]
    assert out["id"].tolist() == [2, 1, 2, 1]
    assert out["500.100000"].tolist() == [4., 2., 24., 22.]
    assert out["600.200000"].tolist() == [8., 4., 208., 204.]
    assert out["x"].tolist() == [20., 10., 20., 10.]
    expected = before.loc[[1, 0, 5, 4]]
    np.testing.assert_array_equal(out[["UMAP_1", "UMAP_2"]].to_numpy(), expected[["UMAP_1", "UMAP_2"]].to_numpy())
    assert out["UMAP cluster"].astype(str).tolist() == expected["Cluster"].tolist()
    assert out["source_file_id"].tolist() == expected["source_file_id"].tolist()
    assert conditions["extra"]["export_selection_written_pixels"] == 4
    pd.testing.assert_frame_equal(before, frames["PCA"])
    assert all(Path(path).read_bytes() == data for path, data in files_before.items())
    assert modern["rds"].read_bytes() == b"immutable RDS fixture"


@pytest.mark.parametrize("all_request", [None, {"mode": "all"}])
@pytest.mark.parametrize("exclude_unused,n", [(False, 9), (True, 6)])
def test_all_mode_preserves_existing_export_and_unused_option(modern, tmp_path, all_request, exclude_unused, n):
    frames = modern["frames"]
    selection = resolve_selection(frames, all_request, SCOPE)
    assert selection is None
    path, name = exp._export_tims(modern["folder"], _lookups(frames), "csv",
        selection=selection, metadata_rds=modern["rds"], metadata_plot=frames["PCA"],
        exclude_unused=exclude_unused, out_dir=tmp_path)
    out = pd.read_csv(path)
    assert name == "UMAP_cluster_TIMS.csv" and len(out) == n
    assert int(out["id"].eq(99).sum()) == (0 if exclude_unused else 3)
    assert set(out["group"].dropna()) == {"Ctrl", "KO1", "KO2"}


def test_selection_runs_before_group_mean_sd_and_n(modern, tmp_path):
    frames = modern["frames"]
    opts = {"categories": ["intensity", "section"], "mode": "group", "group_keys": ["section"]}
    path, name = exp._export_tims(modern["folder"], _lookups(frames), "csv",
        selection=_choose(frames, {"KO1"}), metadata_rds=modern["rds"], metadata_plot=frames["PCA"],
        options=opts, exclude_unused=False, out_dir=tmp_path)
    out = pd.read_csv(path)
    assert "grouped_selected" in name and len(out) == 1
    assert out.loc[0, "group"] == "KO1" and out.loc[0, "n"] == 2
    assert out.loc[0, "500.100000_mean"] == 13.
    assert out.loc[0, "500.100000_sd"] == pytest.approx(np.std([12., 14.], ddof=1))
    assert out.loc[0, "600.200000_mean"] == 106.
    assert out.loc[0, "600.200000_sd"] == pytest.approx(np.std([104., 108.], ddof=1))


def test_selection_survives_identifier_and_intensity_output_disabled(modern, tmp_path):
    frames = modern["frames"]
    path, _ = exp._export_tims(modern["folder"], _lookups(frames), "parquet",
        selection=_choose(frames, {"KO2"}), metadata_rds=modern["rds"], metadata_plot=frames["PCA"],
        options={"categories": ["section", "cluster"]}, out_dir=tmp_path)
    out = pd.read_parquet(path)
    assert len(out) == 2 and set(out["group"]) == {"KO2"}
    assert "id" not in out and "500.100000" not in out
    assert out["UMAP cluster"].astype(str).tolist() == ["4", "3"]


def test_multiple_methods_select_union_and_keep_method_specific_clusters(modern, tmp_path):
    base = modern["frames"]["PCA"]
    frames = OrderedDict(PCA=base.loc[base.group.ne("KO2")].copy(),
                         Harmony=base.loc[base.group.ne("Ctrl")].copy())
    frames["Harmony"]["Cluster"] = ["h1", "h2", "h3", "h4"]
    combined = pd.concat(frames.values()).drop_duplicates(["source_file_id", "source_pixel_id"])
    selection = _choose(frames, {"Ctrl", "KO2"})
    assert selection.summary()["method_counts"] == {"Harmony": 2, "PCA": 2}
    path, _ = exp._export_tims(modern["folder"], _lookups(frames), "csv",
        selection=selection, metadata_rds=modern["rds"], metadata_plot=combined, out_dir=tmp_path)
    out = pd.read_csv(path, keep_default_na=False)
    assert out["group"].tolist() == ["Ctrl", "Ctrl", "KO2", "KO2"]
    assert out["PCA"].astype(str).tolist() == ["2", "1", "", ""]
    assert out["Harmony"].tolist() == ["", "", "h4", "h3"]
    assert out["500.100000"].tolist() == [4., 2., 24., 22.]


@pytest.mark.parametrize("fmt", ["csv", "parquet", "xlsx"])
def test_mz_only_lists_selected_source_features_not_excluded_file_features(modern, tmp_path, fmt):
    for index, entry in enumerate(modern["manifest"]["files"]):
        source = Path(entry["path"])
        raw = pd.read_parquet(source)
        raw[f"{700 + index}.000000"] = 0.  # 全0でも元ファイルのfeature一覧には残す。
        raw.to_parquet(source, index=False)
    frames = modern["frames"]
    conditions = {"extra": {}}
    path, _ = exp._export_tims(modern["folder"], _lookups(frames), fmt,
        selection=_choose(frames, {"Ctrl", "KO2"}), metadata_rds=modern["rds"], metadata_plot=frames["PCA"],
        options={"categories": ["mzlist"]}, conditions=conditions, out_dir=tmp_path)
    out = (pd.read_excel(path, sheet_name=SHEET_NAME, dtype={"列名": str}) if fmt == "xlsx" else
           pd.read_csv(path, dtype={"列名": str}) if fmt == "csv" else pd.read_parquet(path))
    assert set(out["列名"]) == {"500.100000", "600.200000", "700.000000", "702.000000"}
    assert "非ゼロ検出一覧ではない" in conditions["extra"]["mz_list_scope"]


@pytest.mark.parametrize("fmt", ["csv", "parquet", "xlsx"])
def test_legacy_sample_filter_uses_exact_sample_and_preserves_values(tmp_path, fmt):
    raw = pd.DataFrame({"id": [1, 2, 3, 4, 5, 6, 7], "x": [1., 2., 3., 4., 5., 6., 7.],
                        "y": 0., "500.100000": [2., 4., 12., 14., 22., 24., 999.],
                        "annotation": ["Ctrl"] * 2 + ["KO1"] * 2 + ["KO2"] * 2 + ["unused"]})
    source = tmp_path / "mixed.csv"; raw.to_csv(source, index=False)
    frame = pd.DataFrame({"Sample": raw.annotation.iloc[:6], "group": raw.annotation.iloc[:6],
                          "CellID": [f"cell{i}" for i in range(6)], "SpatialX": raw.x.iloc[:6],
                          "SpatialY": 0., "Cluster": ["0", "1", "2", "3", "4", "5"],
                          "UMAP_1": np.arange(6) * .5, "UMAP_2": np.arange(6) * -.25})
    frames = OrderedDict(PCA=frame)
    opts = {"categories": ["id", "coords", "intensity", "section", "umap", "cluster"]}
    path, _ = exp._export_tims(str(tmp_path), _lookups(frames), fmt,
        selection=_choose(frames, {"Ctrl", "KO2"}), options=opts,
        extra_lookups=exp._build_extra_lookups(frame, opts), out_dir=tmp_path / "exports")
    out = _read(path, fmt)
    assert out.id.tolist() == [1, 2, 5, 6] and out["500.100000"].tolist() == [2., 4., 22., 24.]
    assert out.annotation.tolist() == ["Ctrl", "Ctrl", "KO2", "KO2"]
    np.testing.assert_array_equal(out[["UMAP_1", "UMAP_2"]], frame[["UMAP_1", "UMAP_2"]].iloc[[0, 1, 4, 5]])


@pytest.mark.parametrize("fault", ["missing", "duplicate"])
def test_modern_bad_source_identity_fails_without_publishing(modern, tmp_path, fault):
    frames = modern["frames"]
    selection = _choose(frames, {"KO1"})
    source = Path(modern["manifest"]["files"][1]["path"])
    raw = pd.read_parquet(source)
    raw = raw.loc[raw.id.ne(1)] if fault == "missing" else pd.concat([raw, raw.iloc[[0]]], ignore_index=True)
    raw.to_parquet(source, index=False)
    out_dir = tmp_path / "exports"
    with pytest.raises(ValueError, match="画素|対応件数"):
        exp._export_tims(modern["folder"], _lookups(frames), "csv", selection=selection,
            metadata_rds=modern["rds"], metadata_plot=frames["PCA"], out_dir=out_dir)
    assert not list(out_dir.glob("*"))


@pytest.mark.parametrize("missing_group", ["KO1", "Ctrl"])
@pytest.mark.parametrize("mz_only", [False, True])
def test_selected_export_checks_only_selected_source_files(modern, tmp_path, missing_group, mz_only):
    # ★ ver75.1: 未選択KO1の保存先消失はCtrl+KO2出力を妨げず、選択Ctrlの消失は停止する。
    frames = modern["frames"]
    index = ("Ctrl", "KO1", "KO2").index(missing_group)
    Path(modern["manifest"]["files"][index]["path"]).unlink()
    out_dir = tmp_path / "exports"
    options = {"categories": ["mzlist"]} if mz_only else None
    kwargs = dict(selection=_choose(frames, {"Ctrl", "KO2"}),
                  metadata_rds=modern["rds"], metadata_plot=frames["PCA"],
                  options=options, out_dir=out_dir)
    if missing_group == "Ctrl":
        with pytest.raises(ValueError, match="入力.*見つかりません"):
            exp._export_tims(modern["folder"], _lookups(frames), "csv", **kwargs)
        assert not list(out_dir.glob("*"))
    else:
        path, _ = exp._export_tims(modern["folder"], _lookups(frames), "csv", **kwargs)
        if mz_only:
            assert set(pd.read_csv(path, dtype={"列名": str})["列名"]) == {"500.100000", "600.200000"}
        else:
            out = pd.read_csv(path)
            assert out["group"].tolist() == ["Ctrl", "Ctrl", "KO2", "KO2"]
            assert out["500.100000"].tolist() == [4., 2., 24., 22.]


def test_empty_selection_rejected_without_any_output(modern, tmp_path):
    frames = modern["frames"]; catalog = build_catalog(frames, SCOPE)
    with pytest.raises(ValueError, match="1つ以上"):
        resolve_selection(frames, {"mode": "selected", "scope": SCOPE, "signature": catalog["signature"], "ids": []}, SCOPE)
    assert not list(tmp_path.glob("UMAP_cluster*"))


@pytest.fixture
def desi(tmp_path):
    folder = tmp_path / "raw"; folder.mkdir()
    source = folder / "sample.txt"
    header = [["", "", "", "", ""], ["", "", "", "A", "B"], ["", "", "", "1", "2"],
              ["", "", "", "100", "200"], ["", "", "", "10", "20"]]
    data = [[str(pixel), str(pixel * 10), "0", str(pixel * 2), str(pixel * 4)] for pixel in range(1, 8)]
    source.write_text("\n".join("\t".join(row) for row in header + data) + "\n", encoding="utf-8")
    manifest = build_section_manifest([{"path": str(source), "available_rois": [], "roi_role": "region"}])
    fid = manifest["files"][0]["file_id"]
    frame = pd.DataFrame([{"CellID": f"cell{pixel}", "Sample": "sample", "source_file_id": fid,
                          "source_pixel_id": str(pixel), "section_id": f"slice{(pixel - 1) // 2}",
                          "section_display_name": f"Slice{(pixel - 1) // 2}", "subject_id": f"M{(pixel - 1) // 2}",
                          "group": ("Ctrl", "KO1", "KO2")[(pixel - 1) // 2],
                          "SpatialX": pixel * 10., "SpatialY": 0., "Cluster": str(pixel + 10)}
                         for pixel in range(1, 7)])
    rds = tmp_path / "result" / "RDS_Files" / "DESI.rds"; rds.parent.mkdir(parents=True)
    (rds.parent.parent / "section_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return folder, source, header, data, rds, OrderedDict(PCA=frame)


@pytest.mark.parametrize("exclude_unused", [False, True])
def test_desi_selected_writer_keeps_all_header_rows_and_pixel_cluster_alignment(desi, tmp_path, exclude_unused):
    folder, source, header, data, rds, frames = desi
    original = source.read_bytes(); before = frames["PCA"].copy(deep=True)
    conditions = {"extra": {}}
    path, name = exp._export_desi(str(folder), _lookups(frames),
        selection=_choose(frames, {"KO1"}), metadata_rds=rds, metadata_plot=before,
        exclude_unused=exclude_unused, conditions=conditions, out_dir=tmp_path)
    assert name == "UMAP_cluster_DESI_selected.xlsx"
    book = openpyxl.load_workbook(path, data_only=True)
    rows = list(book["sample"].values); labels = rows[0]
    assert [list(row[:5]) for row in rows[:5]] == [[value or None for value in row] for row in header]
    assert [list(row[:5]) for row in rows[5:]] == data[2:4]
    assert [row[labels.index("UMAP cluster")] for row in rows[5:]] == ["13", "14"]
    assert [row[labels.index("group")] for row in rows[5:]] == ["KO1", "KO1"]
    assert conditions["extra"]["export_selection_written_pixels"] == 2
    assert source.read_bytes() == original
    pd.testing.assert_frame_equal(before, frames["PCA"])


def test_desi_all_mode_preserves_existing_rows(desi, tmp_path):
    folder, _, _, data, rds, frames = desi
    path, name = exp._export_desi(str(folder), _lookups(frames), metadata_rds=rds,
        metadata_plot=frames["PCA"], exclude_unused=False, selection=None, out_dir=tmp_path)
    rows = list(openpyxl.load_workbook(path, data_only=True)["sample"].values)
    assert name == "UMAP_cluster_DESI.xlsx"
    assert [list(row[:5]) for row in rows[5:]] == data


def test_desi_partial_missing_selection_does_not_leave_completed_xlsx(desi, tmp_path):
    folder, source, header, data, rds, frames = desi
    selection = _choose(frames, {"KO1"})
    source.write_text("\n".join("\t".join(row) for row in header + data[:3] + data[4:]) + "\n", encoding="utf-8")
    out_dir = tmp_path / "exports"
    with pytest.raises(ValueError, match="対応件数|入力に見つかりません"):
        exp._export_desi(str(folder), _lookups(frames), selection=selection, metadata_rds=rds,
            metadata_plot=frames["PCA"], out_dir=out_dir)
    assert not list(out_dir.glob("*"))


def test_legacy_desi_selection_removes_unselected_sheets_and_keeps_headers(tmp_path):
    folder = tmp_path / "legacy_desi"; folder.mkdir()
    header = [["", "", "", "", ""], ["", "", "", "A", "B"], ["", "", "", "1", "2"],
              ["", "", "", "100", "200"], ["", "", "", "10", "20"]]
    records = []
    for offset, sample in enumerate(("Ctrl", "KO1", "KO2")):
        data = [[str(pixel), str(pixel * 10), "0", str(offset * 10 + pixel), str(pixel * 4)] for pixel in (1, 2)]
        (folder / f"{sample}.txt").write_text("\n".join("\t".join(row) for row in header + data) + "\n", encoding="utf-8")
        records.extend({"CellID": f"{sample}:{pixel}", "Sample": sample, "group": sample,
                        "SpatialX": pixel * 10., "SpatialY": 0., "Cluster": str(offset + pixel)}
                       for pixel in (1, 2))
    frames = OrderedDict(PCA=pd.DataFrame(records))
    path, _ = exp._export_desi(str(folder), _lookups(frames), selection=_choose(frames, {"KO1"}),
                               exclude_unused=False, out_dir=tmp_path)
    book = openpyxl.load_workbook(path, data_only=True)
    assert book.sheetnames == ["KO1"]
    rows = list(book["KO1"].values)
    assert [list(row[:5]) for row in rows[:5]] == [[value or None for value in row] for row in header]
    assert [list(row[:5]) for row in rows[5:]] == [["1", "10", "0", "11", "4"], ["2", "20", "0", "12", "8"]]
    assert [row[-1] for row in rows[5:]] == ["2", "3"]


def test_modern_desi_selected_export_ignores_unrelated_txt_not_in_manifest(desi, tmp_path):
    folder, _, _, data, rds, frames = desi
    # ★ ver75.1: 未選択・未登録TXTをsource IDへ解決しようとして停止してはならない。
    (folder / "unused.txt").write_text("解析とは無関係でDESIとしても不正なファイル", encoding="utf-8")
    path, _ = exp._export_desi(str(folder), _lookups(frames),
        selection=_choose(frames, {"Ctrl", "KO2"}), metadata_rds=rds,
        metadata_plot=frames["PCA"], exclude_unused=False, out_dir=tmp_path)
    book = openpyxl.load_workbook(path, data_only=True)
    assert book.sheetnames == ["sample"]
    rows = list(book["sample"].values)
    assert [list(row[:5]) for row in rows[5:]] == data[:2] + data[4:6]


def test_modern_desi_selected_manifest_resolves_same_basename_across_folders(modern, tmp_path):
    # 選択ファイル一覧を単一フォルダのstem辞書に縮約すると片方の切片が消える。
    manifest = deepcopy(modern["manifest"])
    header = [["", "", "", "", ""], ["", "", "", "A", "B"], ["", "", "", "1", "2"],
              ["", "", "", "100", "200"], ["", "", "", "10", "20"]]
    for entry in manifest["files"]:
        raw = pd.read_parquet(entry["path"])
        source = Path(entry["path"]).with_suffix(".txt")
        data = [[str(int(row.id)), str(row.x), str(row.y), str(row["500.100000"]), str(row["600.200000"])]
                for _, row in raw.iterrows()]
        source.write_text("\n".join("\t".join(row) for row in header + data) + "\n", encoding="utf-8")
        entry["path"] = str(source)
    (modern["rds"].parent.parent / "section_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    frames = modern["frames"]
    path, _ = exp._export_desi(modern["folder"], _lookups(frames),
        selection=_choose(frames, {"Ctrl", "KO2"}), metadata_rds=modern["rds"],
        metadata_plot=frames["PCA"], exclude_unused=False, out_dir=tmp_path)
    book = openpyxl.load_workbook(path, data_only=True)
    assert "Skipped" not in book.sheetnames, "由来IDで全行一致した出力を未一致/空欄と報告してはいけない"
    sheets = [sheet for sheet in book if sheet.title != "Conditions"]
    assert len(sheets) == 2
    written = []
    for sheet in sheets:
        rows = list(sheet.values); labels = rows[0]
        written.extend((row[labels.index("group")], row[labels.index("source_file_id")], float(row[3]),
                        str(row[labels.index("UMAP cluster")])) for row in rows[5:])
    assert [row[0] for row in written] == ["Ctrl", "Ctrl", "KO2", "KO2"]
    assert [row[2] for row in written] == [4., 2., 24., 22.]
    assert [row[3] for row in written] == ["2", "1", "4", "3"]
    assert len({row[1] for row in written}) == 2
